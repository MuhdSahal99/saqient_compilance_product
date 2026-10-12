"""Synthetic demonstration content.

Everything here is invented for demonstration. It contains no real
organisations, people, regulations or statutory citations.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

SYNTHETIC_NOTICE = "Synthetic demonstration content — not legal guidance."

POLICY_TITLE = "Vendor Data Handling Policy (synthetic)"

POLICY_SECTIONS: list[tuple[str, list[str]]] = [
    (
        "Vendor Data Handling Policy",
        [
            "This is a synthetic document created for software demonstration. It describes a "
            "fictional company, Example Trading Co., and is not a real policy.",
            "Policy owner: Head of Information Security (fictional role).",
            "Effective date: 01/03/2026. Next review date: 01/03/2027.",
        ],
    ),
    (
        "1. Purpose",
        [
            "This policy sets expectations for how vendors handle data shared by Example Trading Co.",
        ],
    ),
    (
        "2. Data retention",
        [
            "Vendors keep shared data only for as long as the contract requires. Retention periods "
            "are listed in each data sharing schedule.",
            "At the end of the contract, vendors return or delete data within 30 days and confirm "
            "deletion in writing.",
        ],
    ),
    (
        "3. Security incidents",
        [
            "Vendors must notify Example Trading Co. of any security incident affecting shared data "
            "without undue delay.",
            "Our systems are 100% secure, so incidents are not expected.",
        ],
    ),
    (
        "4. Subcontractors",
        [
            "Vendors may use subcontractors only with prior written approval.",
        ],
    ),
]

POLICY_AR_LINES = [
    "سياسة تجريبية — محتوى توضيحي اصطناعي وليس إرشاداً قانونياً.",
    "مالك السياسة: رئيس أمن المعلومات (دور خيالي).",
]


def policy_text() -> str:
    parts: list[str] = []
    for heading, paras in POLICY_SECTIONS:
        parts.append(heading)
        parts.extend(paras)
    return "\n\n".join(parts) + "\n"


def policy_txt_bytes() -> bytes:
    return policy_text().encode("utf-8")


def policy_docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    for i, (heading, paras) in enumerate(POLICY_SECTIONS):
        d.add_heading(heading, level=0 if i == 0 else 1)
        for p in paras:
            d.add_paragraph(p)
    for line in POLICY_AR_LINES:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def policy_pdf_bytes(blank_last_page: bool = False) -> bytes:
    """Two-page PDF with real extractable text, page 2 starting at section 3."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    def write_sections(sections: list[tuple[str, list[str]]]) -> None:
        y = height - 72
        for heading, paras in sections:
            c.setFont("Helvetica-Bold", 13)
            c.drawString(72, y, heading)
            y -= 26
            c.setFont("Helvetica", 10)
            for para in paras:
                line = ""
                for word in para.split():
                    if c.stringWidth(line + " " + word, "Helvetica", 10) > width - 144:
                        c.drawString(72, y, line.strip())
                        y -= 14
                        line = ""
                    line += " " + word
                c.drawString(72, y, line.strip())
                y -= 28  # blank line between paragraphs

    write_sections(POLICY_SECTIONS[:3])
    c.showPage()
    write_sections(POLICY_SECTIONS[3:])
    c.showPage()
    if blank_last_page:
        c.showPage()
    c.save()
    return buf.getvalue()


def blank_pdf_bytes() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(100, 100, 200, 200)  # a drawing, but no text: looks like a scanned page
    c.showPage()
    c.save()
    return buf.getvalue()


CHECKLIST_NAME = "Policy Baseline Checklist (synthetic demo)"
CHECKLIST_DESCRIPTION = (
    "A small synthetic checklist used to demonstrate the review workflow. " + SYNTHETIC_NOTICE
)

DEMO_SOURCE = {
    "title": "Saqient synthetic demonstration source",
    "publisher": "Saqient (synthetic)",
    "version_label": "demo-1",
    "excerpt": SYNTHETIC_NOTICE,
}

DEMO_REQUIREMENTS: list[dict[str, object]] = [
    {
        "code": "DEMO-01",
        "title": "Policy owner is named",
        "category": "Governance",
        "severity": "medium",
        "description": "The policy should name an accountable owner.",
        "human_instructions": "Confirm the named owner is a real role in the organisation.",
        "check_config": {
            "type": "required_text",
            "phrases": ["policy owner", "document owner"],
            "match": "any",
        },
    },
    {
        "code": "DEMO-02",
        "title": "Effective date is stated",
        "category": "Governance",
        "severity": "low",
        "description": "The policy should state when it takes effect.",
        "human_instructions": "Check the date is plausible and consistent with approvals.",
        "check_config": {"type": "pattern", "pattern": "effective date {any} {date}"},
    },
    {
        "code": "DEMO-03",
        "title": "Data retention is addressed",
        "category": "Data management",
        "severity": "high",
        "description": "The policy should describe how long data is kept.",
        "human_instructions": "Check retention periods are specific, not open-ended.",
        "check_config": {"type": "required_text", "phrases": ["retention"]},
    },
    {
        "code": "DEMO-04",
        "title": "Incident notification is described",
        "category": "Incidents",
        "severity": "high",
        "description": "The policy should require incidents to be notified.",
        "human_instructions": "Check a notification timeframe or trigger is defined.",
        "check_config": {
            "type": "phrase_groups",
            "groups": [["incident"], ["notify", "notification", "report"]],
        },
    },
    {
        "code": "DEMO-05",
        "title": "No absolute assurance language",
        "category": "Language",
        "severity": "medium",
        "description": "Policies should avoid absolute claims about security or compliance.",
        "human_instructions": "Review flagged wording in context; it may be quoted or negated.",
        "check_config": {
            "type": "prohibited_terms",
            "terms": ["100% secure", "guaranteed compliance", "fully compliant"],
        },
    },
    {
        "code": "DEMO-06",
        "title": "Periodic access review is required",
        "category": "Access",
        "severity": "high",
        "description": "The policy should require periodic review of vendor access.",
        "human_instructions": "Look for equivalent wording the check may have missed.",
        "check_config": {
            "type": "required_text",
            "phrases": ["access review", "review of access"],
            "match": "any",
        },
    },
    {
        "code": "DEMO-07",
        "title": "Document is classified as a policy",
        "category": "Metadata",
        "severity": "low",
        "description": "The uploaded document's type metadata should be 'policy'.",
        "human_instructions": "Correct the document type if it was mis-labelled.",
        "check_config": {"type": "metadata", "field": "document_type", "equals": "policy"},
    },
    {
        "code": "DEMO-08",
        "title": "Management approval is evidenced",
        "category": "Governance",
        "severity": "medium",
        "description": "There should be evidence the policy was approved by management.",
        "human_instructions": "Ask for approval minutes or a signed approval record.",
        "check_config": {"type": "manual"},
    },
]


def write_samples(out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "synthetic-vendor-policy.txt").write_bytes(policy_txt_bytes())
    (out_dir / "synthetic-vendor-policy.docx").write_bytes(policy_docx_bytes())
    (out_dir / "synthetic-vendor-policy.pdf").write_bytes(policy_pdf_bytes())
    (out_dir / "README.md").write_text(
        "# Sample files\n\nSynthetic documents for trying the workflow. "
        + SYNTHETIC_NOTICE
        + "\nRegenerate with `uv run python -m app.demo_content <dir>`.\n"
    )


if __name__ == "__main__":
    write_samples(Path(sys.argv[1] if len(sys.argv) > 1 else "samples"))
